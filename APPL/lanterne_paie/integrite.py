"""Contrôle d'intégrité : la copie de travail conserve onglets, formules et extensions de l'original."""

from __future__ import annotations

import re
import warnings
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from .modeles import Champ, Statut
from .plan_excel import Ecriture

SOURCE = "Contrôle d'intégrité"


def _formules_et_valeurs(chemin: Path) -> tuple[list[str], dict[tuple[str, str], str], dict[tuple[str, str], object]]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(chemin, data_only=False)
    formules, valeurs = {}, {}
    try:
        for ws in wb.worksheets:
            for ligne in ws.iter_rows():
                for c in ligne:
                    v = c.value
                    if v is None:
                        continue
                    texte = v.text if hasattr(v, "text") else v  # formules matricielles
                    if isinstance(texte, str) and texte.startswith("="):
                        formules[(ws.title, c.coordinate)] = _normaliser_formule(texte)
                    else:
                        valeurs[(ws.title, c.coordinate)] = v
        return wb.sheetnames, formules, valeurs
    finally:
        wb.close()


def _normaliser_formule(f: str) -> str:
    """Écriture canonique : espaces, casse, préfixe _xlfn et FALSE()/TRUE() équivalents à FALSE/TRUE."""
    f = re.sub(r"\s+", "", f).upper().replace("_XLFN.", "")
    return f.replace("FALSE()", "FALSE").replace("TRUE()", "TRUE")


def _extensions(chemin: Path) -> dict[str, int]:
    """Nombre de validations de données et de mises en forme conditionnelles par feuille (XML brut)."""
    resultat = {}
    with zipfile.ZipFile(chemin) as z:
        for nom in z.namelist():
            if nom.startswith("xl/worksheets/sheet") and nom.endswith(".xml"):
                xml = z.read(nom).decode("utf-8", "replace")
                resultat[nom] = len(re.findall(r"<(?:x14:)?dataValidation[ >]", xml)) + len(
                    re.findall(r"<(?:x14:)?cfRule[ >]", xml))
    return resultat


def verifier(original: Path, copie: Path, ecritures: list[Ecriture], extensions: bool = True) -> list[Champ]:
    onglets_o, formules_o, valeurs_o = _formules_et_valeurs(original)
    onglets_c, formules_c, valeurs_c = _formules_et_valeurs(copie)
    champs = []

    manquants = [o for o in onglets_o if o not in onglets_c]
    champs.append(Champ("integrite_onglets", "Onglets conservés", f"{len(onglets_c)} onglets", SOURCE,
                        Statut.ERREUR if manquants else Statut.VALIDE,
                        f"Onglets manquants : {', '.join(manquants)}" if manquants else "", True, modifiable=False))

    modifiees = [k for k, f in formules_o.items() if formules_c.get(k) != f]
    detail = ", ".join(f"{f}!{c}" for f, c in modifiees[:10])
    champs.append(Champ("integrite_formules", "Formules conservées", f"{len(formules_o)} formules", SOURCE,
                        Statut.ERREUR if modifiees else Statut.VALIDE,
                        f"{len(modifiees)} formule(s) modifiée(s) ou supprimée(s) : {detail}" if modifiees else "",
                        True, modifiable=False))

    prevues = {(e.feuille, e.cellule) for e in ecritures}
    imprevues = [k for k in set(valeurs_o) | set(valeurs_c)
                 if k not in prevues and k not in formules_o and k not in formules_c and _valeur(valeurs_o.get(k)) != _valeur(valeurs_c.get(k))]
    detail = ", ".join(f"{f}!{c}" for f, c in sorted(imprevues)[:10])
    champs.append(Champ("integrite_saisies", "Seules les cellules prévues sont modifiées", f"{len(prevues)} cellule(s) prévue(s)",
                        SOURCE, Statut.ERREUR if imprevues else Statut.VALIDE,
                        f"Modifications imprévues : {detail}" if imprevues else "", True, modifiable=False))

    if extensions:
        ext_o, ext_c = _extensions(original), _extensions(copie)
        perdues = [n for n, nb in ext_o.items() if ext_c.get(n, 0) < nb]
        champs.append(Champ("integrite_extensions", "Listes déroulantes et mises en forme conditionnelles",
                            f"{sum(ext_c.values())} / {sum(ext_o.values())}", SOURCE,
                            Statut.ERREUR if perdues else Statut.VALIDE,
                            f"Éléments perdus dans : {', '.join(perdues)}" if perdues else "", True, modifiable=False))
    return champs


def _valeur(v: object) -> object:
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.lstrip("'")
    return v
