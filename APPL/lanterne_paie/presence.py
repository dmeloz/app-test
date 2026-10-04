"""Lecture de la fiche de présence mensuelle (PDF natif ou classeur Excel)."""

from __future__ import annotations

import re
from pathlib import Path

from .controles import normaliser
from .dates_fr import lire_date
from .modeles import FichePresence, Intervenant
from .pdf_texte import Segment, rangees, segments

FONCTIONS = ["Savant·e", "Naïf·ve", "Artiste", "Musicien·ne"]
ROLES_ORGANISATION = ["Responsable", "Comptabilité"]
_LIBELLES = {
    "seance": "Séance",
    "programme": "Programme",
    "cinema": "Cinéma",
    "animation": "Animation",
    "organisation": "Organisation",
    "infos": "Informations importantes",
}
_TEL = re.compile(r"(?:\+41\s?|0)\d{2}\s?\d{3}\s?\d{2}\s?\d{2}")
_EMAIL = re.compile(r"[^\s,;]+@[^\s,;]+\.[A-Za-z]{2,}")


def lire_presence(chemin: Path) -> FichePresence:
    if chemin.suffix.lower() == ".pdf":
        lignes = _lignes_pdf(segments(chemin))
    elif chemin.suffix.lower() in {".xlsx", ".xlsm"}:
        lignes = _lignes_excel(chemin)
    else:
        raise ValueError(f"Format de fiche de présence non pris en charge : {chemin.name}")
    return analyser_lignes(chemin, lignes)


def _lignes_pdf(segs: list[Segment]) -> list[tuple[str, list[str]]]:
    """Chaque rangée devient (libellé de gauche ou "", valeurs de droite)."""
    resultat = []
    rangs = rangees(segs)
    if not rangs:
        return resultat
    marge = min(s.x0 for s in segs)
    for rang in rangs:
        gauche = [s.texte for s in rang if s.x0 < marge + 150]
        droite = [s.texte for s in rang if s.x0 >= marge + 150]
        resultat.append((" ".join(gauche), droite))
    return resultat


def _lignes_excel(chemin: Path) -> list[tuple[str, list[str]]]:
    """Même représentation à partir d'un classeur : 1re cellule non vide = libellé."""
    from openpyxl import load_workbook

    classeur = load_workbook(chemin, read_only=True, data_only=True)
    resultat = []
    try:
        for feuille in classeur.worksheets:
            for ligne in feuille.iter_rows(values_only=True):
                cellules = [str(v).strip() for v in ligne if v is not None and str(v).strip()]
                if not cellules:
                    continue
                if cellules[0] in FONCTIONS or cellules[0] in ROLES_ORGANISATION or cellules[0] in _LIBELLES.values():
                    resultat.append((cellules[0], cellules[1:]))
                else:
                    resultat.append(("", cellules))
    finally:
        classeur.close()
    return resultat


def analyser_lignes(chemin: Path, lignes: list[tuple[str, list[str]]]) -> FichePresence:
    fiche = FichePresence(source=chemin)
    section = ""
    fonction_courante: Intervenant | None = None
    for i, (libelle, valeurs) in enumerate(lignes):
        texte = " ".join([libelle, *valeurs]).strip()
        if "CONFIRMATION DE PR" in texte.upper() and not fiche.club:
            fiche.club = re.split(r"CONFIRMATION", texte, flags=re.I)[0].strip()
            continue
        if libelle:
            section = libelle
            fonction_courante = None
        if section == _LIBELLES["seance"]:
            for v in valeurs:
                m_date = re.search(r"\d{1,2}\.\d{1,2}\.\d{2,4}", v)
                m_heure = re.search(r"\b\d{1,2}:\d{2}\b", v)
                if m_date and not fiche.date_seance:
                    fiche.date_seance = lire_date(m_date.group(0))
                if m_heure and not fiche.heure:
                    fiche.heure = m_heure.group(0)
        elif section == _LIBELLES["programme"] and valeurs and not fiche.programme:
            fiche.programme = valeurs[0]
        elif section == _LIBELLES["cinema"] and valeurs and not fiche.cinema:
            fiche.cinema = valeurs[0]
        elif section in ROLES_ORGANISATION:
            if libelle and valeurs:
                fiche.organisation[section] = _intervenant(section, valeurs)
            elif section in fiche.organisation and valeurs:
                _completer_contacts(fiche.organisation[section], " ".join(valeurs))
        elif section in FONCTIONS:
            if libelle and valeurs:
                fonction_courante = _intervenant(section, valeurs)
                fiche.intervenants.append(fonction_courante)
            elif fonction_courante is not None and valeurs:
                _completer_contacts(fonction_courante, " ".join(valeurs))
        if "cachet de l" in normaliser(texte) or (section and "CHF" in texte and not fiche.cachet_artiste):
            m = re.search(r"CHF\s*([\d'’.]+)", texte)
            if m is None and i + 1 < len(lignes):
                m = re.search(r"CHF\s*([\d'’.]+)", " ".join([lignes[i + 1][0], *lignes[i + 1][1]]))
            if m and "cachet" in normaliser(texte + " ".join(lignes[max(0, i - 1)][1])):
                fiche.cachet_artiste = m.group(1).rstrip(".")
    return fiche


def _intervenant(fonction: str, valeurs: list[str]) -> Intervenant:
    """« Prénom Nom , Rue 1 , 1071 Localité » (+ éventuels contacts sur la même ligne)."""
    texte = " ".join(valeurs)
    personne = Intervenant(fonction=fonction, nom_complet="")
    _completer_contacts(personne, texte)
    texte = _EMAIL.sub(" ", texte)
    texte = _TEL.sub(" ", texte)
    parties = [p.strip() for p in texte.split(",") if p.strip()]
    if parties:
        personne.nom_complet = re.sub(r"\s+", " ", parties[0])
    if len(parties) >= 3:
        personne.adresse = re.sub(r"\s+", " ", parties[1])
        personne.npa_localite = re.sub(r"\s+", " ", " ".join(parties[2:]))
    elif len(parties) == 2:
        personne.adresse = re.sub(r"\s+", " ", parties[1])
    return personne


def _completer_contacts(personne: Intervenant, texte: str) -> None:
    for tel in _TEL.findall(texte):
        tel = re.sub(r"\s+", " ", tel.strip())
        if tel not in personne.telephones and not re.fullmatch(r"\d{4}", tel):
            personne.telephones.append(tel)
    m = _EMAIL.search(texte)
    if m and not personne.email:
        personne.email = m.group(0)
