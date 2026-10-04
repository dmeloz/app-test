"""Lecture SEULE du classeur « Gestion administrative travailleurs ».

Le classeur n'est jamais enregistré par cette bibliothèque : openpyxl supprimerait les listes
déroulantes et mises en forme conditionnelles (extensions Excel 2010). Toute écriture passe par
Microsoft Excel (voir excel_mac.py).
"""

from __future__ import annotations

import warnings
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .controles import normaliser
from .modeles import Collaborateur

PREMIERE_LIGNE_COLLAB = 6
DERNIERE_LIGNE_COLLAB = 40
LIGNE_PREMIERE_SEANCE = 24  # Configuration!B24 : séance 4 de la saison précédente
LIGNE_DERNIERE_SEANCE = 35
LIGNES_PAR_BLOC = 11
PLACES_PAR_BLOC = 8

ENTETES_COLLAB = {
    "A": "Prénom", "B": "Nom", "C": "Collaborateur·trice", "D": "Adresse", "E": "NPA, Localité",
    "F": "E-mail", "G": "Téléphone", "H": "Statut", "I": "Taux LPP", "J": "Déclaration AVS",
    "K": "Date naiss.", "L": "N° AVS", "M": "IBAN",
}
# Colonnes de saisie (la colonne C est une formule et n'est jamais écrite).
CHAMPS_COLLAB = {
    "prenom": "A", "nom": "B", "adresse": "D", "npa_localite": "E", "email": "F",
    "telephone": "G", "statut": "H", "taux_lpp": "I", "declaration_avs": "J",
    "date_naissance": "K", "no_avs": "L", "iban": "M",
}


class ClasseurInvalide(Exception):
    pass


@dataclass
class Feuilles:
    configuration: str
    collaborateurs: str
    prestations: str
    fiche: str


@dataclass
class LignePrestation:
    ligne: int
    collaborateur: str
    fonction: str


@dataclass
class Classeur:
    chemin: Path
    feuilles: Feuilles
    annee: int | None
    raison_sociale: str
    adresse_club: list[str]
    seances: dict[int, date | None]  # ligne de Configuration -> date
    libelles_seances: dict[int, str]
    tarifs: dict[str, tuple[float | None, float | None]]  # fonction -> (salarié, indépendant)
    collaborateurs: list[Collaborateur]
    prestations: dict[int, list[LignePrestation]]  # ligne de Configuration -> lignes saisies
    anomalies: list[str] = field(default_factory=list)

    # --- correspondances séance / lignes -----------------------------------------------------
    @staticmethod
    def ligne_configuration(numero_seance: int) -> int:
        """Séances 1 à 3 de la saison : B30:B32 ; séances 4 à 9 : B24:B29 (classeur en année civile)."""
        if not 1 <= numero_seance <= 9:
            raise ValueError(f"Numéro de séance hors plage : {numero_seance}")
        return 29 + numero_seance if numero_seance <= 3 else 20 + numero_seance

    @staticmethod
    def lignes_bloc(ligne_config: int) -> range:
        """Lignes de saisie du bloc Prestations correspondant à une ligne de Configuration."""
        k = ligne_config - LIGNE_PREMIERE_SEANCE
        debut = 7 + LIGNES_PAR_BLOC * k
        return range(debut, debut + PLACES_PAR_BLOC)

    def ligne_de_date(self, jour: date) -> int | None:
        for ligne, valeur in self.seances.items():
            if valeur == jour:
                return ligne
        return None

    def rechercher(self, nom_complet: str = "", nom: str = "", prenom: str = "", no_avs: str = "") -> list[Collaborateur]:
        """Correspondances par N° AVS, sinon par « Prénom Nom » (ou nom + prénom)."""
        if no_avs:
            par_avs = [c for c in self.collaborateurs if c.no_avs.strip() == no_avs.strip()]
            if par_avs:
                return par_avs
        cibles = set()
        if nom_complet:
            cibles.add(normaliser(nom_complet))
        if nom or prenom:
            cibles.add(normaliser(f"{prenom} {nom}"))
            cibles.add(normaliser(f"{nom} {prenom}"))
        return [
            c for c in self.collaborateurs
            if normaliser(c.cle) in cibles or normaliser(f"{c.nom} {c.prenom}") in cibles
        ]

    def premiere_ligne_collab_libre(self) -> int | None:
        occupees = {c.ligne for c in self.collaborateurs}
        for ligne in range(PREMIERE_LIGNE_COLLAB, DERNIERE_LIGNE_COLLAB + 1):
            if ligne not in occupees:
                return ligne
        return None


def detecter_feuilles(noms: list[str]) -> Feuilles:
    """Repère les onglets par leur numéro et un libellé approché, sans nom codé en dur."""

    def trouver(numero: str, mot: str) -> str:
        for nom in noms:
            n = normaliser(nom)
            if n.startswith(f"{numero}.") and mot in n:
                return nom
        for nom in noms:
            if mot in normaliser(nom):
                return nom
        raise ClasseurInvalide(f"Onglet « {numero}. {mot} » introuvable (onglets : {', '.join(noms)})")

    return Feuilles(
        configuration=trouver("1", "configuration"),
        collaborateurs=trouver("2", "collaborat"),
        prestations=trouver("3", "prestation"),
        fiche=trouver("4", "fiche"),
    )


def _texte(valeur: object) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).strip()


def _jour(valeur: object) -> date | None:
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    return None


def lire_classeur(chemin: Path) -> Classeur:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # extensions non gérées par openpyxl : lecture seule
        valeurs = load_workbook(chemin, data_only=True)
        formules = load_workbook(chemin, data_only=False)
    try:
        feuilles = detecter_feuilles(valeurs.sheetnames)
        conf = valeurs[feuilles.configuration]
        collab = valeurs[feuilles.collaborateurs]
        prest = valeurs[feuilles.prestations]
        anomalies = _verifier_structure(formules, feuilles)

        annee = conf["B5"].value if isinstance(conf["B5"].value, int) else None
        if annee is None:
            anomalies.append("Configuration!B5 (année civile) n'est pas une année.")
        seances = {r: _jour(conf[f"B{r}"].value) for r in range(LIGNE_PREMIERE_SEANCE, LIGNE_DERNIERE_SEANCE + 1)}
        libelles = {r: _texte(conf[f"A{r}"].value) for r in seances}
        tarifs = {}
        for r in range(13, 22):
            fonction = _texte(conf[f"A{r}"].value)
            if fonction:
                tarifs[fonction] = (conf[f"B{r}"].value, conf[f"C{r}"].value)

        collaborateurs = []
        for r in range(PREMIERE_LIGNE_COLLAB, DERNIERE_LIGNE_COLLAB + 1):
            prenom, nom = _texte(collab[f"A{r}"].value), _texte(collab[f"B{r}"].value)
            if not prenom and not nom:
                continue
            lpp = collab[f"I{r}"].value
            collaborateurs.append(Collaborateur(
                ligne=r, prenom=prenom, nom=nom,
                adresse=_texte(collab[f"D{r}"].value), npa_localite=_texte(collab[f"E{r}"].value),
                email=_texte(collab[f"F{r}"].value), telephone=_texte(collab[f"G{r}"].value),
                statut=_texte(collab[f"H{r}"].value),
                taux_lpp=float(lpp) if isinstance(lpp, (int, float)) else None,
                declaration_avs=_texte(collab[f"J{r}"].value),
                date_naissance=_jour(collab[f"K{r}"].value),
                no_avs=_texte(collab[f"L{r}"].value), iban=_texte(collab[f"M{r}"].value),
            ))

        prestations = {}
        for r in seances:
            lignes = []
            for ligne in Classeur.lignes_bloc(r):
                nom_c = _texte(prest[f"D{ligne}"].value)
                if nom_c:
                    lignes.append(LignePrestation(ligne, nom_c, _texte(prest[f"E{ligne}"].value)))
            prestations[r] = lignes

        return Classeur(
            chemin=chemin, feuilles=feuilles, annee=annee,
            raison_sociale=_texte(conf["B8"].value),
            adresse_club=[x.strip(" ,") for x in (_texte(conf["B9"].value) + "\n" + _texte(conf["B10"].value)).splitlines() if x.strip(" ,")],
            seances=seances, libelles_seances=libelles, tarifs=tarifs,
            collaborateurs=collaborateurs, prestations=prestations, anomalies=anomalies,
        )
    finally:
        valeurs.close()
        formules.close()


def _verifier_structure(classeur, feuilles: Feuilles) -> list[str]:
    """Contrôle que les cellules utilisées sont bien à l'emplacement attendu."""
    anomalies = []
    collab: Worksheet = classeur[feuilles.collaborateurs]
    for col, attendu in ENTETES_COLLAB.items():
        trouve = _texte(collab[f"{col}5"].value)
        if normaliser(trouve) != normaliser(attendu):
            anomalies.append(f"Collaborateurs!{col}5 : en-tête « {trouve} » au lieu de « {attendu} ».")
    for r in range(PREMIERE_LIGNE_COLLAB, DERNIERE_LIGNE_COLLAB + 1):
        if not str(collab[f"C{r}"].value or "").startswith("="):
            anomalies.append(f"Collaborateurs!C{r} n'est plus une formule.")
            break
    prest: Worksheet = classeur[feuilles.prestations]
    if normaliser(_texte(prest["D6"].value)) != normaliser("Collaborateur·trice"):
        anomalies.append("Prestations!D6 : en-tête « Collaborateur·trice » introuvable.")
    fiche: Worksheet = classeur[feuilles.fiche]
    conf: Worksheet = classeur[feuilles.configuration]
    # Les feuilles sont protégées : seules les cellules de saisie déverrouillées peuvent être écrites.
    cibles = [(collab, f"{c}{PREMIERE_LIGNE_COLLAB}") for c in CHAMPS_COLLAB.values()]
    cibles += [(conf, f"B{r}") for r in range(LIGNE_PREMIERE_SEANCE, LIGNE_DERNIERE_SEANCE + 1)]
    cibles += [(prest, f"{c}{r}") for c in "DE" for r in Classeur.lignes_bloc(LIGNE_PREMIERE_SEANCE)]
    cibles += [(fiche, "B5"), (fiche, "E5")]
    verrouillees = [f"{ws.title}!{ref}" for ws, ref in cibles if ws.protection.sheet and ws[ref].protection.locked]
    if verrouillees:
        anomalies.append("Cellules de saisie verrouillées (protection) : " + ", ".join(verrouillees[:5]))
    if not str(fiche["D8"].value or "").startswith("="):
        anomalies.append("Fiche de salaire!D8 (contrôle) n'est pas une formule.")
    if not str(fiche["F44"].value or "").startswith("="):
        anomalies.append("Fiche de salaire!F44 (total versé) n'est pas une formule.")
    return anomalies


def annee_du_classeur(chemin: Path) -> int | None:
    """Lecture rapide de Configuration!B5 pour la détection automatique des classeurs."""
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(chemin, read_only=True, data_only=True)
        try:
            feuilles = detecter_feuilles(wb.sheetnames)
            valeur = wb[feuilles.configuration]["B5"].value
            return valeur if isinstance(valeur, int) else None
        finally:
            wb.close()
    except Exception:
        return None
